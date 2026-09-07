import importlib.util
import tempfile
import unittest
from pathlib import Path


DEPENDENCIES_AVAILABLE = all(
    importlib.util.find_spec(name) is not None for name in ("openpyxl", "pandas")
)


@unittest.skipUnless(DEPENDENCIES_AVAILABLE, "workbook dependencies are not installed")
class WorkbookToolTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Font

        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ID", "Group", "Value", "Department"])
        sheet.append([1, "Alpha", "A", "Nursing"])
        sheet.append([1, "Alpha", "B", "Nursing"])
        sheet.append([2, "alpha", "C", "Research"])
        sheet["A2"].font = Font(bold=True)
        sheet["A2"].comment = Comment("Keep me", "CEIR")
        workbook.save(self.source)
        workbook.close()

    def tearDown(self):
        self.temp.cleanup()

    def test_tab_splitter_uses_case_insensitive_unique_sheet_names(self):
        from openpyxl import load_workbook
        from tools.tab_splitter import run_tab_splitter

        output = self.root / "tabs.xlsx"
        run_tab_splitter(
            {"input_file": str(self.source), "output_file": str(output), "header_row": 1, "split_column": "Group"},
            lambda _message: None,
        )
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Alpha", "alpha_1"])
        finally:
            workbook.close()

    def test_dynamic_pivot_groups_by_identifier(self):
        import pandas as pd
        from tools.dynamic_pivot import run_dynamic_pivot

        output = self.root / "pivot.xlsx"
        run_dynamic_pivot(
            {
                "input_file": str(self.source),
                "output_file": str(output),
                "source_sheet": "Sheet",
                "header_row": 1,
                "key_column": "ID",
                "pivot_column": "Value",
            },
            lambda _message: None,
        )
        result = pd.read_excel(output)
        first = result[result["ID"] == 1].iloc[0]
        self.assertEqual(first["Value 1"], "A")
        self.assertEqual(first["Value 2"], "B")

    def test_dynamic_pivot_rejects_conflicting_carry_values(self):
        from openpyxl import load_workbook
        from tools.dynamic_pivot import run_dynamic_pivot

        workbook = load_workbook(self.source)
        workbook.active["D3"] = "Different"
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "Department"):
            run_dynamic_pivot(
                {
                    "input_file": str(self.source),
                    "output_file": str(self.root / "pivot.xlsx"),
                    "source_sheet": "Sheet",
                    "header_row": 1,
                    "key_column": "ID",
                    "pivot_column": "Value",
                },
                lambda _message: None,
            )

    def test_failed_transform_preserves_existing_output(self):
        from tools.tab_splitter import run_tab_splitter

        output = self.root / "existing.xlsx"
        output.write_bytes(b"existing")
        with self.assertRaises(ValueError):
            run_tab_splitter(
                {"input_file": str(self.source), "output_file": str(output), "header_row": 1, "split_column": "Missing"},
                lambda _message: None,
            )
        self.assertEqual(output.read_bytes(), b"existing")

    def test_workbook_splitter_stages_batch_and_preserves_basic_cell_details(self):
        from openpyxl import load_workbook
        from tools.workbook_splitter import run_workbook_splitter

        output = self.root / "output"
        run_workbook_splitter(
            {
                "input_file": str(self.source),
                "output_folder": str(output),
                "header_row": 1,
                "split_column": "Group",
                "tab_column": None,
            },
            lambda _message: None,
        )
        generated = sorted(output.glob("*.xlsx"))
        self.assertEqual(len(generated), 2)
        self.assertEqual(len({path.name.casefold() for path in generated}), 2)
        alpha = load_workbook(next(path for path in generated if "Alpha" in path.name))
        try:
            self.assertTrue(alpha.active["A2"].font.bold)
            self.assertEqual(alpha.active["A2"].comment.text, "Keep me")
        finally:
            alpha.close()

    def test_cancelled_batch_split_publishes_no_workbooks(self):
        from core.file_utils import OperationCancelled
        from tools.workbook_splitter import run_workbook_splitter

        output = self.root / "output"
        with self.assertRaises(OperationCancelled):
            run_workbook_splitter(
                {
                    "input_file": str(self.source),
                    "output_folder": str(output),
                    "header_row": 1,
                    "split_column": "Group",
                    "tab_column": None,
                },
                lambda _message: None,
                lambda: True,
            )
        self.assertEqual(list(output.glob("*.xlsx")), [])

    def test_duplicate_headers_are_rejected_consistently(self):
        from openpyxl import load_workbook
        from core.excel_helpers import read_active_sheet_headers

        workbook = load_workbook(self.source)
        workbook.active["D1"] = "id"
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "Duplicate header"):
            read_active_sheet_headers(self.source, 1)


if __name__ == "__main__":
    unittest.main()
