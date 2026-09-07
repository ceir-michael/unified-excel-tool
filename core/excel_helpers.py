from copy import copy
from pathlib import Path

from openpyxl import load_workbook


def copy_cell(source, target) -> None:
    """Copy a cell value and its supported formatting to another cell."""
    target.value = source.value

    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)

    if source.comment:
        target.comment = copy(source.comment)

    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)


def find_header_column(ws, header_name: str, header_row: int = 1) -> int | None:
    """Return the column index containing a header, or None if not found."""
    wanted = header_name.strip().casefold()

    for column in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=column).value
        if value is not None and str(value).strip().casefold() == wanted:
            return column

    return None


def worksheet_headers(ws, header_row: int = 1) -> dict[str, int]:
    """Return unambiguous worksheet headers and reject duplicates."""
    headers: dict[str, int] = {}
    folded: dict[str, str] = {}
    for column in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=column).value
        if value is None or not str(value).strip():
            continue
        name = str(value).strip()
        key = name.casefold()
        if key in folded:
            raise ValueError(
                f"Duplicate header '{name}' was found. Rename duplicate columns before continuing."
            )
        folded[key] = name
        headers[name] = column
    return headers


def read_active_sheet_headers(path: Path, header_row: int) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(worksheet_headers(workbook.active, header_row))
    finally:
        workbook.close()
